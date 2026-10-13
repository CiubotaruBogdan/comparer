import tempfile
import unittest
from pathlib import Path

import fitz
from docx import Document

from document_core import (
    ComparisonCancelled,
    compare_documents,
    read_document,
    read_docx,
    read_pdf,
)
from notifications import UserFacingError


class DocumentCoreTests(unittest.TestCase):
    def _make_docx_with_stories_and_tables(self, path: Path) -> None:
        document = Document()
        document.add_paragraph("Paragraf din corp")

        outer = document.add_table(rows=1, cols=1)
        outer.cell(0, 0).text = "Celulă exterioară"
        nested = outer.cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Celulă imbricată"

        horizontal = document.add_table(rows=1, cols=2)
        horizontal.cell(0, 0).merge(horizontal.cell(0, 1)).text = "Celulă fuzionată orizontal"

        vertical = document.add_table(rows=2, cols=1)
        vertical.cell(0, 0).merge(vertical.cell(1, 0)).text = "Celulă fuzionată vertical"

        repeated_rows = document.add_table(rows=3, cols=1)
        for index, row in enumerate(repeated_rows.rows):
            row.cells[0].text = f"Celulă separată {index}"

        first_section = document.sections[0]
        first_section.header.paragraphs[0].text = "Antet identic"
        first_header_table = first_section.header.add_table(rows=1, cols=1, width=1000000)
        first_header_table.cell(0, 0).text = "Tabel în antet"
        first_header_nested = first_header_table.cell(0, 0).add_table(rows=1, cols=1)
        first_header_nested.cell(0, 0).text = "Tabel imbricat în antet"
        first_section.footer.paragraphs[0].text = "Antet identic"
        first_footer_table = first_section.footer.add_table(rows=1, cols=1, width=1000000)
        first_footer_table.cell(0, 0).text = "Tabel în subsol"

        second_section = document.add_section()
        second_section.header.is_linked_to_previous = False
        second_section.header.paragraphs[0].text = "Antet identic"
        second_section.footer.is_linked_to_previous = False
        second_section.footer.paragraphs[0].text = "Subsol secțiunea doi"
        document.save(path)

    def test_docx_reads_header_footer_tables_nested_and_merged_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "structurat.docx"
            self._make_docx_with_stories_and_tables(path)
            texts = read_docx(str(path))

        self.assertIn("[HEADER] Tabel în antet", texts)
        self.assertIn("[HEADER] Tabel imbricat în antet", texts)
        self.assertIn("[FOOTER] Tabel în subsol", texts)
        # Same visible words in distinct header parts and in a footer are valid
        # distinct stories; de-duplication must be by part identity plus role.
        self.assertEqual(texts.count("[HEADER] Antet identic"), 2)
        self.assertEqual(texts.count("[FOOTER] Antet identic"), 1)
        self.assertIn("Paragraf din corp", texts)
        self.assertIn("Celulă exterioară", texts)
        self.assertIn("Celulă imbricată", texts)
        self.assertEqual(texts.count("Celulă fuzionată orizontal"), 1)
        self.assertEqual(texts.count("Celulă fuzionată vertical"), 1)
        self.assertEqual(
            [text for text in texts if text.startswith("Celulă separată")],
            ["Celulă separată 0", "Celulă separată 1", "Celulă separată 2"],
        )

    def test_normalized_matching_retains_original_display_text_and_result_options(self):
        left = ["  Salut   LUME  ", "[HEADER] Antet vechi", "Text corp"]
        right = ["salut lume", "[HEADER] Antet nou", "Text corp"]
        supplied_options = {
            "ignore_whitespace": True,
            "ignore_case": True,
            "ignore_headers": True,
        }

        result = compare_documents(left, right, supplied_options)

        self.assertEqual(result["options"], supplied_options)
        self.assertEqual(result["similarity"], 1.0)
        self.assertEqual(result["stats"], {
            "total_left": 2,
            "total_right": 2,
            "equal": 2,
            "modified": 0,
            "added": 0,
            "deleted": 0,
        })
        self.assertEqual(result["diff_count"], 0)
        self.assertEqual(result["left_blocks"][0], ("equal", "  Salut   LUME  ", False))
        self.assertEqual(result["right_blocks"][0], ("equal", "salut lume", False))
        self.assertNotIn("[HEADER] Antet vechi", [block[1] for block in result["left_blocks"]])

    def test_pdf_encryption_and_image_only_documents_have_localized_explanations(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            encrypted_path = directory_path / "parolat.pdf"
            encrypted = fitz.open()
            encrypted.new_page()
            encrypted.save(
                encrypted_path,
                encryption=fitz.PDF_ENCRYPT_AES_256,
                owner_pw="proprietar",
                user_pw="secret",
            )
            encrypted.close()

            with self.assertRaisesRegex(UserFacingError, "protejat prin parolă"):
                read_pdf(str(encrypted_path))

            scanned_path = directory_path / "scanat.pdf"
            scanned = fitz.open()
            scanned.new_page()
            scanned.save(scanned_path)
            scanned.close()
            with self.assertRaisesRegex(UserFacingError, "scanat.*OCR"):
                read_document(str(scanned_path))

            empty_path = directory_path / "gol.pdf"
            empty_path.write_bytes(b"")
            with self.assertRaisesRegex(UserFacingError, "PDF este gol"):
                read_pdf(str(empty_path))

    def test_cancellation_is_checked_by_readers_and_comparison(self):
        with tempfile.TemporaryDirectory() as directory:
            docx_path = Path(directory) / "anulare.docx"
            document = Document()
            for index in range(10):
                document.add_paragraph(f"Paragraf {index}")
            document.save(docx_path)

            calls = 0

            def cancel_after_several_checks():
                nonlocal calls
                calls += 1
                return calls >= 5

            with self.assertRaises(ComparisonCancelled):
                read_docx(str(docx_path), cancel_after_several_checks)
            self.assertGreaterEqual(calls, 5)

        with self.assertRaises(ComparisonCancelled):
            compare_documents(["stânga"] * 10, ["dreapta"] * 10, cancel=lambda: True)

        with self.assertRaises(ComparisonCancelled):
            read_pdf("oricare.pdf", cancel=lambda: True)


if __name__ == "__main__":
    unittest.main()
